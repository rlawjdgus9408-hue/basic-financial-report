#!/usr/bin/env python3
"""기초재무분석 엑셀 형식 검사.

RAW 고정 형식(references/raw-format.md)과 Graph 차트 제목·범례 규칙
(references/ratios-and-charts.md 「차트 제목·범례 = 원천 표」)을 검사한다.
이상이 하나라도 있으면 종료 코드 1.

사용: python scripts/check_workbook.py 파일.xlsx
"""
import sys
from openpyxl import load_workbook

SHEETS = ["Graph", "Analysis", "RAW"]
GUBUN = {"재무상태표", "손익계산서", "검증"}
ROWTYPE = {"합계", "소계", "세부", "차감", "검증"}
VERIFY_ROWS = ["대차 차이", "원본 자산총계", "자산총계 대사 차이", "원본 당기순이익", "당기순이익 대사 차이"]
HEADER_ROW = 5


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def check_raw(wb, errs, warns):
    if "RAW" not in wb.sheetnames:
        errs.append("[R1] 'RAW' 시트 없음")
        return
    ws = wb["RAW"]
    if "단위: 원" not in str(ws["A2"].value or ""):
        errs.append(f"[R2] A2 단위 표기가 '(단위: 원)'이 아님: {ws['A2'].value!r}")
    if not str(ws["A3"].value or "").startswith("출처:"):
        errs.append(f"[R2] A3 출처 표기 없음: {ws['A3'].value!r}")

    # 머리글
    hdr = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    while hdr and hdr[-1] in (None, ""):
        hdr.pop()
    if hdr[:3] != ["구분", "계정과목", "행유형"]:
        errs.append(f"[R3] 5행 A~C 머리글이 '구분·계정과목·행유형'이 아님: {hdr[:3]}")
        return
    if not hdr or hdr[-1] != "비고":
        errs.append(f"[R3] 마지막 머리글이 '비고'가 아님: {hdr[-1] if hdr else None!r}")
        return
    years_raw = hdr[3:-1]
    years = []
    for y in years_raw:
        try:
            years.append(int(y))
        except (TypeError, ValueError):
            errs.append(f"[R3] 연도 머리글이 숫자 연도가 아님: {y!r}")
            return
    if not years:
        errs.append("[R3] 연도 열 없음")
        return
    if years != sorted(years) or len(set(years)) != len(years):
        errs.append(f"[R3] 연도가 오래된 순(좌→우)이 아니거나 중복: {years}")
    ycols = list(range(4, 4 + len(years)))

    # 병합·숨김·그룹
    if ws.merged_cells.ranges:
        errs.append(f"[R4] 병합 셀 {len(ws.merged_cells.ranges)}개: {[str(r) for r in ws.merged_cells.ranges][:5]}")
    hidden = [r for r, d in ws.row_dimensions.items() if d.hidden]
    grouped = [r for r, d in ws.row_dimensions.items() if d.outlineLevel]
    hcols = [k for k, d in ws.column_dimensions.items() if d.hidden]
    if hidden or grouped or hcols:
        errs.append(f"[R4] 숨김·그룹 존재 (숨김행 {hidden[:5]}, 그룹행 {grouped[:5]}, 숨김열 {hcols[:5]})")
    if ws.freeze_panes != "D6":
        warns.append(f"[R4] 틀 고정이 D6이 아님: {ws.freeze_panes}")

    # 행 검사
    order, blocks_seen = [], []
    verify_names = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        a, b, c = (ws.cell(r, k).value for k in (1, 2, 3))
        vals = [ws.cell(r, k).value for k in ycols]
        if a in (None, "") and b in (None, "") and c in (None, ""):
            if any(v not in (None, "") for v in vals):
                errs.append(f"[R5] {r}행: 구분·계정·행유형 없이 숫자만 있음")
            continue
        if a not in GUBUN:
            errs.append(f"[R5] {r}행 구분 값 오류: {a!r}")
        if c not in ROWTYPE:
            errs.append(f"[R5] {r}행 행유형 값 오류: {c!r}")
        if not blocks_seen or blocks_seen[-1] != a:
            blocks_seen.append(a)
        if a == "검증":
            verify_names.append(str(b).strip())
            continue
        for k, v in zip(ycols, vals):
            ref = ws.cell(r, k).coordinate
            if c in ("합계", "소계") and v not in (None, "") and not is_formula(v):
                errs.append(f"[R6] {ref} {b}: {c} 행인데 숫자 직접 입력")
            if c in ("합계", "소계") and v in (None, ""):
                warns.append(f"[R6] {ref} {b}: {c} 행인데 비어 있음")
            if c in ("세부", "차감") and is_formula(v):
                errs.append(f"[R6] {ref} {b}: {c} 행인데 수식 (입력값이어야 함)")
            if c in ("세부", "차감") and isinstance(v, float) and not v.is_integer():
                errs.append(f"[R6] {ref} {b}: 원 단위 정수가 아님 ({v})")
            if c == "차감" and isinstance(v, (int, float)) and v < 0:
                errs.append(f"[R7] {ref} {b}: 차감 계정이 음수 (양수 입력 후 상위 수식에서 차감)")

    if blocks_seen != ["재무상태표", "손익계산서", "검증"]:
        errs.append(f"[R8] 블록 순서가 재무상태표 → 손익계산서 → 검증이 아니거나 블록이 섞임: {blocks_seen}")
    missing = [n for n in VERIFY_ROWS if n not in verify_names]
    if missing:
        errs.append(f"[R9] 검증 블록 행 누락: {missing}")


def _title_text(title):
    if title is None or title.tx is None:
        return None
    tx = title.tx
    if tx.strRef is not None:
        return ("REF", tx.strRef.f)
    if tx.rich is not None:
        parts = []
        for p in tx.rich.p or []:
            for run in (p.r or []):
                parts.append(run.t or "")
        return "".join(parts)
    return None


def _resolve(wb, ref):
    ref = ref.replace("$", "")
    sh, cell = ref.rsplit("!", 1)
    sh = sh.strip("'")
    if sh not in wb.sheetnames:
        return None
    return wb[sh][cell.split(":")[0]].value


FORBIDDEN = {"PieChart": "파이", "PieChart3D": "파이(3D)", "DoughnutChart": "도넛",
             "AreaChart": "영역", "AreaChart3D": "영역(3D)", "BarChart3D": "3D 막대",
             "LineChart3D": "3D 선", "SurfaceChart3D": "3D 표면"}


def check_graph(wb, errs, warns):
    if "Graph" not in wb.sheetnames:
        errs.append("[G0] 'Graph' 시트 없음")
        return
    ws = wb["Graph"]
    texts = {str(c.value).strip() for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and not c.value.startswith("=")}
    charts = ws._charts
    if len(charts) != 5:
        warns.append(f"[G0] 차트 {len(charts)}개 (기준 5개)")
    for i, ch in enumerate(charts, 1):
        name = type(ch).__name__
        if name in FORBIDDEN:
            errs.append(f"[G3] 차트{i}: 사용 금지 유형 ({FORBIDDEN[name]})")
        t = _title_text(ch.title)
        if t is None or t == "":
            errs.append(f"[G1] 차트{i}: 제목 없음")
        else:
            tval = _resolve(wb, t[1]) if isinstance(t, tuple) else t
            if tval is None or str(tval).strip() not in texts:
                errs.append(f"[G1] 차트{i}: 제목 '{tval}'이 Graph 시트의 표 제목 셀과 일치하지 않음")
        sers = list(ch.series)
        for sub in getattr(ch, "_charts", []):  # 콤보 차트
            sers += list(sub.series)
        for j, s in enumerate(sers, 1):
            tx = s.tx
            if tx is None or tx.strRef is None:
                lit = getattr(tx, "v", None) if tx is not None else None
                errs.append(f"[G2] 차트{i} 계열{j}: 범례 이름이 셀 참조가 아님 (직접 입력값 {lit!r})")
                continue
            v = _resolve(wb, tx.strRef.f)
            if v in (None, ""):
                errs.append(f"[G2] 차트{i} 계열{j}: 범례 참조 셀 {tx.strRef.f}이 비어 있음")


def main(path):
    wb = load_workbook(path)
    errs, warns = [], []
    if wb.sheetnames != SHEETS:
        warns.append(f"[S0] 시트 구성·순서가 Graph·Analysis·RAW가 아님: {wb.sheetnames}")
    check_raw(wb, errs, warns)
    check_graph(wb, errs, warns)
    for w in warns:
        print("WARN ", w)
    for e in errs:
        print("ERROR", e)
    print(f"\n결과: 이상 {len(errs)}건, 경고 {len(warns)}건" + (" — 정상" if not errs else ""))
    return 1 if errs else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
