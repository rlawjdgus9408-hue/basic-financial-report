#!/usr/bin/env python3
"""check_workbook.py 회귀 테스트. 정상본은 이상 0건, 위반본은 지정한 이상이 모두 검출돼야 한다.
사용: python tests/test_check_workbook.py
"""
import os, sys, subprocess, tempfile
from openpyxl import Workbook
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.chart.series import SeriesLabel, StrRef

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "..", "scripts", "check_workbook.py")

BS = [("자산총계", "합계", "=D7+D10"), ("I. 유동자산", "소계", "=D8+D9"),
      ("현금및현금성자산", "세부", 100), ("매출채권", "세부", 50),
      ("II. 비유동자산", "소계", "=D11-D12"), ("기계장치", "세부", 80),
      ("감가상각누계액", "차감", 20), ("부채총계", "합계", "=D14"),
      ("단기차입금", "세부", 60), ("자본총계", "합계", "=D16+D17"),
      ("자본금", "세부", 50), ("이익잉여금", "세부", 100)]


def build(path, broken=False):
    wb = Workbook()
    g = wb.active; g.title = "Graph"
    wb.create_sheet("Analysis"); r = wb.create_sheet("RAW")
    r["A1"] = "㈜테스트 표준재무제표 RAW"; r["A2"] = "(단위: 원)"
    r["A3"] = "출처: 표준재무제표증명, 2026-10-01, 원본 단위 원"
    years = [2023, 2024, 2025] if not broken else [2025, 2024, 2023]
    r.append([]) if False else None
    for i, h in enumerate(["구분", "계정과목", "행유형"] + years + ["비고"], 1):
        r.cell(5, i, h)
    row = 6
    for name, typ, v in BS:
        r.cell(row, 1, "재무상태표"); r.cell(row, 2, name); r.cell(row, 3, typ)
        for k in range(3):
            col = "DEF"[k]
            val = v.replace("D", col) if isinstance(v, str) else v
            if broken and name == "자산총계":
                val = 210
            if broken and name == "감가상각누계액":
                val = -20
            r.cell(row, 4 + k, val)
        row += 1
    row += 1
    for name, typ, v in [("매출액", "세부", 500), ("매출원가", "세부", 300), ("매출총이익", "소계", None)]:
        r.cell(row, 1, "손익계산서"); r.cell(row, 2, name); r.cell(row, 3, typ)
        for k in range(3):
            col = "DEF"[k]
            r.cell(row, 4 + k, v if v is not None else f"={col}{row-2}-{col}{row-1}")
        row += 1
    row += 1
    vnames = ["대차 차이", "원본 자산총계", "자산총계 대사 차이", "원본 당기순이익", "당기순이익 대사 차이"]
    if broken:
        vnames = vnames[:3]
    for n in vnames:
        r.cell(row, 1, "검증"); r.cell(row, 2, n); r.cell(row, 3, "검증"); row += 1
    r.freeze_panes = "D6"
    if broken:
        r.merge_cells("A1:C1")

    # Graph: 원천 표 + 차트
    g["B2"] = "수익성 지표"
    for i, y in enumerate([2023, 2024, 2025]):
        g.cell(3, 3 + i, y)
    g["B4"] = "원가율"; g["B5"] = "영업이익률"
    for i in range(3):
        g.cell(4, 3 + i, 0.6); g.cell(5, 3 + i, 0.1)
    ch = PieChart() if broken else LineChart()
    ch.title = "수익성 지표" if not broken else "영업이익률 유지 메커니즘"
    for rr in (4, 5):
        data = Reference(g, min_col=3, max_col=5, min_row=rr)
        ch.add_data(data, from_rows=True, titles_from_data=False)
        s = ch.series[-1]
        s.tx = SeriesLabel(v="원가율") if broken else SeriesLabel(strRef=StrRef(f"'Graph'!$B${rr}"))
    g.add_chart(ch, "G2")
    wb.save(path)


def run(path):
    p = subprocess.run([sys.executable, SCRIPT, path], capture_output=True, text=True)
    return p.returncode, p.stdout


def main():
    ok = True
    with tempfile.TemporaryDirectory() as d:
        good, bad = os.path.join(d, "good.xlsx"), os.path.join(d, "bad.xlsx")
        build(good); build(bad, broken=True)
        code, out = run(good)
        if code != 0 or "ERROR" in out:
            ok = False; print("정상본에서 이상 검출:\n" + out)
        code, out = run(bad)
        for tag in ["[R3]", "[R4]", "[R6]", "[R7]", "[R9]", "[G1]", "[G2]", "[G3]"]:
            if tag not in out:
                ok = False; print(f"위반본에서 {tag} 미검출")
        if code != 1:
            ok = False; print("위반본 종료 코드가 1이 아님")
        if not ok:
            print(out)
    print("전체 통과" if ok else "실패")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
