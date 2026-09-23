"""
raw_input.json -> RAW 시트 엑셀 작성 스크립트.
.claude/skills/raw-sheet/SKILL.md §3, §4를 구현한다.

사용법:
    python scripts/build_raw.py raw_input.json out.xlsx
    python scripts/build_raw.py raw_input.json 기존.xlsx --into

--into: output_xlsx가 이미 존재하는 워크북일 때, 그 안의 'RAW' 시트에만 값을 채운다
        (다른 시트·서식은 건드리지 않는다). 시트가 없으면 새로 만든다.
        'RAW' 시트에 이미 내용(값이 있는 셀)이 있으면 덮어쓰지 않고 중단한다.

작성 후 결과 파일을 다시 읽어 JSON 값과 한 칸씩 대조해 자동 검증한다(§4 절차 4).
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

_NUM_FORMAT = r'#,##0;(#,##0);"-"'
_FONT_NAME = "맑은 고딕"
_BORDER = Border(bottom=Side(style="thin"))

# 블록 순서 (§3-1). JSON에 이 목록에 없는 title이 있으면 뒤에 이어 붙인다.
_TITLE_ORDER = ["재무상태표", "손익계산서", "제조원가명세서", "이익잉여금처분계산서", "결손금처리계산서"]


def _sheet_has_content(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                return True
    return False


def _write_statement(ws, start_row, title, years, rows):
    """블록 하나(제목행 + 헤더행 + 데이터행들)를 쓰고, (다음 블록 시작행, 이 블록의 헤더행)을 반환."""
    r = start_row
    header_cell = ws.cell(row=r, column=2, value=f"{title}           (단위: 원)")
    header_cell.font = Font(name=_FONT_NAME, size=18, bold=True)
    header_cell.alignment = Alignment(horizontal="right")
    header_cell.border = _BORDER
    ws.row_dimensions[r].height = 28.5
    r += 1

    year_header_row = r
    label_header = ws.cell(row=r, column=2, value="계정과목")
    label_header.font = Font(name=_FONT_NAME, size=10, bold=True)
    label_header.border = _BORDER
    for idx, year in enumerate(years):
        c = ws.cell(row=r, column=3 + idx, value=year)
        c.font = Font(name=_FONT_NAME, size=10, bold=True)
        c.alignment = Alignment(horizontal="center")
        c.border = _BORDER
    r += 1

    for row in rows:
        level = row.get("level", 3)
        is_bold = level in (0, 1)
        label_cell = ws.cell(row=r, column=2, value=row.get("name", ""))
        label_cell.font = Font(name=_FONT_NAME, size=10, bold=is_bold)
        label_cell.alignment = Alignment(indent=level)
        label_cell.border = _BORDER

        values = row.get("values")
        for idx in range(len(years)):
            v = values[idx] if values is not None and idx < len(values) else None
            c = ws.cell(row=r, column=3 + idx, value=v if v is not None else 0)
            c.font = Font(name=_FONT_NAME, size=10, bold=is_bold)
            c.number_format = _NUM_FORMAT
            c.alignment = Alignment(horizontal="right")
            c.border = _BORDER
        r += 1

    return r + 2, year_header_row  # 블록 사이 빈 행 2개 (§3-1)


def build_workbook(data, ws=None):
    """RAW 시트를 채운다. ws가 주어지면 그 시트에, 없으면 새 워크북을 만들어 채운다."""
    years = data.get("years") or []
    if not years:
        raise ValueError("years 배열이 비어 있습니다.")

    if ws is None:
        wb = openpyxl.Workbook()
        ws = wb.active
    else:
        wb = ws.parent
    ws.title = "RAW"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for idx in range(len(years)):
        ws.column_dimensions[get_column_letter(3 + idx)].width = 16

    statements = {s.get("title"): s for s in data.get("statements", [])}
    ordered_titles = list(_TITLE_ORDER) + [t for t in statements if t not in _TITLE_ORDER]

    row_cursor = 2
    first_header_row = None
    for title in ordered_titles:
        stmt = statements.get(title)
        if not stmt:
            continue
        row_cursor, header_row = _write_statement(ws, row_cursor, title, years, stmt.get("rows", []))
        if first_header_row is None:
            first_header_row = header_row

    if first_header_row is not None:
        # 틀 고정: 첫 블록의 헤더 행 아래, B열 오른쪽 (§3-5)
        ws.freeze_panes = f"C{first_header_row + 1}"

    return wb, ws


def _matched_title(label, titles):
    for t in titles:
        if t and label.startswith(t):
            return t
    return None


def verify(output_path, data):
    """작성된 엑셀을 다시 읽어 JSON 값과 한 칸씩 대조한다 (§4 절차 4)."""
    wb = openpyxl.load_workbook(output_path, data_only=True)
    ws = wb["RAW"]

    titles = list(_TITLE_ORDER) + [s.get("title") for s in data.get("statements", [])]
    section_title = None
    year_row_seen = False
    rows_by_section = {}

    for r in range(1, ws.max_row + 1):
        b = ws.cell(row=r, column=2).value
        label = str(b).strip() if b is not None else ""
        if not label:
            continue
        matched = _matched_title(label, titles)
        if matched:
            section_title = matched
            rows_by_section.setdefault(section_title, [])
            year_row_seen = False
            continue
        if section_title is None:
            continue
        if label == "계정과목":
            year_row_seen = True
            continue
        if not year_row_seen:
            continue
        rows_by_section[section_title].append((r, label))

    mismatches = []
    for stmt in data.get("statements", []):
        title = stmt.get("title")
        sheet_rows = rows_by_section.get(title, [])
        json_rows = stmt.get("rows", [])
        if len(sheet_rows) != len(json_rows):
            mismatches.append(f"[{title}] 행 수 불일치: 시트 {len(sheet_rows)}개 vs JSON {len(json_rows)}개")
            continue
        for (r, label), row in zip(sheet_rows, json_rows):
            expected_name = row.get("name", "")
            if label != expected_name:
                mismatches.append(f"[{title}] row{r}: 계정명 불일치 '{label}' != '{expected_name}'")
            values = row.get("values") or []
            for idx, expected in enumerate(values):
                actual = ws.cell(row=r, column=3 + idx).value
                expected_cmp = 0 if expected is None else expected
                actual_cmp = 0 if actual is None else actual
                if abs(float(actual_cmp) - float(expected_cmp)) > 1:
                    mismatches.append(
                        f"[{title}] '{label}' {idx + 1}번째 연도: 시트값 {actual_cmp} != JSON값 {expected_cmp}"
                    )
    return mismatches


def main():
    parser = argparse.ArgumentParser(description="raw_input.json -> RAW 시트 엑셀 작성")
    parser.add_argument("input_json")
    parser.add_argument("output_xlsx")
    parser.add_argument("--into", action="store_true", help="output_xlsx가 기존 워크북이면 그 RAW 시트만 채운다")
    args = parser.parse_args()

    data = json.loads(Path(args.input_json).read_text(encoding="utf-8"))

    if args.into:
        out_path = Path(args.output_xlsx)
        if not out_path.exists():
            print(f"오류: --into 옵션인데 '{out_path}' 파일이 없습니다.")
            sys.exit(2)
        wb = openpyxl.load_workbook(out_path)
        if "RAW" in wb.sheetnames:
            existing_ws = wb["RAW"]
            if _sheet_has_content(existing_ws):
                print("오류: 'RAW' 시트에 이미 내용이 있습니다. 덮어쓰지 않고 중단합니다.")
                sys.exit(1)
            ws = existing_ws
        else:
            ws = wb.create_sheet("RAW")
        wb, ws = build_workbook(data, ws=ws)
    else:
        wb, ws = build_workbook(data)

    wb.save(args.output_xlsx)
    print(f"작성 완료: {args.output_xlsx}")

    mismatches = verify(args.output_xlsx, data)
    if mismatches:
        print(f"\n[검증 실패] {len(mismatches)}건 — 작성된 엑셀이 JSON 값과 다릅니다:")
        for m in mismatches[:50]:
            print(f"  x {m}")
        sys.exit(1)

    print("[검증 통과] 작성된 엑셀이 JSON 값과 한 칸씩 일치합니다.")


if __name__ == "__main__":
    main()
