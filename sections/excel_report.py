"""
기초재무분석 엑셀(Graph · Analysis · RAW) 생성 — 기초재무진단 스킬 「엑셀 산출물 형식」.

skills/기초재무진단/references/excel-output.md, raw-format.md, ratios-and-charts.md(Q4-4)를 따른다.
Word 보고서와 같은 자동분석 결과(session_state의 계정·지표 코멘트)를 쓰므로 두 산출물의 숫자와
코멘트가 같다. 만든 파일은 스킬의 scripts/check_workbook.py로 형식을 검사할 수 있다.

RAW의 합계·소계는 수식으로 쓰되, 원본 번호 체계(Ⅰ. / (1) / 1. / 가.)로 추정한 하위 행의 합이
원본 값과 실제로 같을 때만 수식으로 바꾼다. 같지 않으면 원본 값을 그대로 입력하고 비고에 남긴다 —
수식 추정이 틀려도 숫자가 원본과 달라지는 일은 없다.
"""
import datetime
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_FONT = "맑은 고딕"
_NUM_FMT = "#,##0;-#,##0"
_BLUE = "0000FF"
_HEADER_ROW = 5
_FIRST_YEAR_COL = 4  # RAW D열부터 연도

_ROMAN_RE = re.compile(r"^\s*(?:[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+|[IVX]+)\s*\.")
_CONTRA_RE = re.compile(r"충당금|누계액|정부보조금|할인발행차금|현재가치할인")
_RATIO_ROW_RE = re.compile(r"률|%|/")


def _font(bold=False, color=None, size=10, italic=False):
    return Font(name=_FONT, size=size, bold=bold, color=color, italic=italic)


# ── RAW 계층 추정 ─────────────────────────────────────────────────────────────

def _level(label, prev_level):
    """원본 번호 체계로 계층을 추정한다. 0=총계, 1=Ⅰ., 2=(1), 3=1., 4=가.
    번호가 없는 행(예: "(대손충당금)", "예치금")은 바로 앞 세부 행과 같은 계층으로 본다."""
    s = str(label).strip()
    if "총계" in s or "합계" in s:
        return 0
    if _ROMAN_RE.match(s):
        return 1
    if re.match(r"^\(\d+\)", s):
        return 2
    if re.match(r"^\d+\.", s):
        return 3
    if re.match(r"^[가-하]\.", s):
        return 4
    if prev_level is None:
        return 3
    return prev_level if prev_level >= 2 else prev_level + 1


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if f != f else f  # NaN → None


def _statement_rows(df, years):
    rows, prev = [], None
    for _, r in df.iterrows():
        label = str(r["계정과목"]).strip()
        lvl = _level(label, prev)
        prev = lvl
        contra = bool(_CONTRA_RE.search(label))
        values = [_num(r.get(y)) for y in years]
        if contra:
            values = [abs(v) if v is not None else None for v in values]
        rows.append({"label": label, "level": lvl, "contra": contra, "values": values,
                     "children": None, "signs": None, "note": ""})
    return rows


def _sums_match(parent, children, signs, n_years):
    for k in range(n_years):
        total = sum(s * (c["values"][k] or 0) for c, s in zip(children, signs))
        if abs(total - (parent["values"][k] or 0)) > 1:
            return False
    return True


def _assign_children(rows, n_years, is_income):
    """각 행의 하위 행과 부호를 정하고, 하위 합이 원본 값과 같을 때만 수식 대상으로 표시한다."""
    for i, row in enumerate(rows):
        if row["level"] == 0:
            continue
        span = []
        for j in range(i + 1, len(rows)):
            if rows[j]["level"] <= row["level"]:
                break
            span.append(j)
        if span:
            child_level = min(rows[j]["level"] for j in span)
            kids = [rows[j] for j in span if rows[j]["level"] == child_level]
            signs = [-1 if k["contra"] else 1 for k in kids]
            if _sums_match(row, kids, signs, n_years):
                row["children"], row["signs"] = kids, signs
            else:
                row["note"] = "하위 계정 합과 원본 금액이 달라 원본 값 입력"

    top = [r for r in rows if r["level"] == 1]
    if is_income:
        # 이익 행(매출총이익·영업이익·차감전순이익·당기순이익) = 앞 이익 행 ± 그 사이 수익·비용 행.
        anchor, between = None, []
        for r in top:
            is_profit = re.search(r"이익|손익|손실", r["label"]) and "영업외" not in r["label"]
            if anchor is not None and is_profit and r["children"] is None:
                kids = [anchor] + between
                signs = [1] + [1 if re.search(r"수익|이익", k["label"]) else -1 for k in between]
                if _sums_match(r, kids, signs, n_years):
                    r["children"], r["signs"], r["note"] = kids, signs, ""
                anchor, between = r, []
            elif anchor is None:
                anchor = r
            else:
                between.append(r)
    else:
        # 자산총계 = 앞쪽 Ⅰ·Ⅱ 합, 부채총계·자본총계도 같은 방식. "부채와 자본총계" = 부채총계 + 자본총계.
        last_total, block_l1, totals_since_assets = None, [], []
        for r in rows:
            if r["level"] == 1:
                block_l1.append(r)
            elif r["level"] == 0:
                candidates = [(block_l1, [1] * len(block_l1)),
                              (totals_since_assets, [1] * len(totals_since_assets))]
                for kids, signs in candidates:
                    if kids and _sums_match(r, kids, signs, n_years):
                        r["children"], r["signs"] = kids, signs
                        break
                if "자산" in r["label"] and "부채" not in r["label"]:
                    totals_since_assets = []
                else:
                    totals_since_assets.append(r)
                block_l1 = []


# ── RAW 시트 ──────────────────────────────────────────────────────────────────

def _write_raw(ws, company, years, bs_rows, is_rows):
    ycols = [get_column_letter(_FIRST_YEAR_COL + k) for k in range(len(years))]
    note_col = _FIRST_YEAR_COL + len(years)

    ws["A1"] = f"{company} 표준재무제표 RAW"
    ws["A2"] = "(단위: 원)"
    ws["A3"] = (f"출처: 고객 업로드 재무제표(AI 추출), {datetime.date.today():%Y-%m-%d} 수령, 원본 단위 원")
    for k, h in enumerate(["구분", "계정과목", "행유형"] + [int(y) for y in years] + ["비고"], start=1):
        c = ws.cell(_HEADER_ROW, k, h)
        c.font = _font(bold=True)
        c.border = Border(bottom=Side(style="thin"))

    r = _HEADER_ROW + 1
    for block, rows in (("재무상태표", bs_rows), ("손익계산서", is_rows)):
        for row in rows:
            row["raw_row"] = r
            r += 1
        r += 1  # 블록 사이 빈 행

    for block, rows in (("재무상태표", bs_rows), ("손익계산서", is_rows)):
        for row in rows:
            rr = row["raw_row"]
            if row["children"]:
                kind = "합계" if row["level"] <= 1 else "소계"
            else:
                kind = "차감" if row["contra"] else "세부"
            ws.cell(rr, 1, block)
            ws.cell(rr, 2, row["label"]).alignment = Alignment(indent=max(row["level"] - 1, 0))
            ws.cell(rr, 3, kind)
            blank_years = []
            for k, col in enumerate(ycols):
                cell = ws[f"{col}{rr}"]
                if row["children"]:
                    terms = "".join(("+" if s > 0 else "-") + f"{col}{c['raw_row']}"
                                    for c, s in zip(row["children"], row["signs"]))
                    cell.value = "=" + terms.lstrip("+")
                    cell.font = _font()
                else:
                    v = row["values"][k]
                    if v is None:
                        blank_years.append(years[k])
                    else:
                        cell.value = int(round(v))
                    cell.font = _font(color=_BLUE)
                cell.number_format = _NUM_FMT
            notes = [row["note"]] if row["note"] else []
            if blank_years:
                notes.append("원본 공란")
            if notes:
                ws.cell(rr, note_col, ", ".join(notes))
            for k in (1, 3):
                ws.cell(rr, k).font = _font()
            ws.cell(rr, 2).font = _font(bold=row["level"] <= 1)

    # 검증 블록
    def _find(rows, pattern, exclude=None):
        for row in rows:
            if re.search(pattern, row["label"]) and not (exclude and re.search(exclude, row["label"])):
                return row
        return None

    assets = _find(bs_rows, r"자산총계|자산합계")
    liab = _find(bs_rows, r"부채총계|부채합계", r"자본")
    equity = _find(bs_rows, r"자본총계|자본합계", r"부채")
    net = _find([x for x in is_rows if x["level"] <= 1], r"당기순")
    vr = r
    verify = [
        ("대차 차이", lambda col: f"={col}{assets['raw_row']}-({col}{liab['raw_row']}+{col}{equity['raw_row']})"
         if assets and liab and equity else None),
        ("원본 자산총계", lambda col, k=None: None),
        ("자산총계 대사 차이", lambda col: f"={col}{assets['raw_row']}-{col}{vr + 1}" if assets else None),
        ("원본 당기순이익", lambda col, k=None: None),
        ("당기순이익 대사 차이", lambda col: f"={col}{net['raw_row']}-{col}{vr + 3}" if net else None),
    ]
    for i, (name, formula) in enumerate(verify):
        rr = vr + i
        ws.cell(rr, 1, "검증").font = _font()
        ws.cell(rr, 2, name).font = _font(bold=True)
        ws.cell(rr, 3, "검증").font = _font()
        for k, col in enumerate(ycols):
            cell = ws[f"{col}{rr}"]
            if name == "원본 자산총계" and assets:
                cell.value, cell.font = int(round(assets["values"][k] or 0)), _font(color=_BLUE)
            elif name == "원본 당기순이익" and net:
                cell.value, cell.font = int(round(net["values"][k] or 0)), _font(color=_BLUE)
            else:
                cell.value, cell.font = formula(col), _font()
            cell.number_format = _NUM_FMT
    ws.cell(vr, note_col, "0이어야 한다").font = _font()

    for col, width in zip("ABC", (10, 30, 8)):
        ws.column_dimensions[col].width = width
    for col in ycols:
        ws.column_dimensions[col].width = 16
    ws.column_dimensions[get_column_letter(note_col)].width = 30
    ws.freeze_panes = "D6"
    for cell in ("A1", "A2", "A3"):
        ws[cell].font = _font(bold=cell == "A1")


# ── Analysis 시트 ─────────────────────────────────────────────────────────────

_LEGEND = [("DAEEF3", "하늘색 — 강점·긍정 (Strength)"), ("F2F2F2", "회색 — 중립·현상 (Neutral)"),
           ("FDE9D9", "주황 — 경고·위험 (Warning)")]


def _write_analysis(ws, company, years, blocks, notes):
    """blocks: [(제목, rows, 계정 코멘트 dict, 지표 제목, 지표 rows, 지표 코멘트 dict)] — 손익계산서 먼저."""
    ycols = [get_column_letter(3 + k) for k in range(len(years))]
    level_col, comment_col = 3 + len(years), 4 + len(years)

    ws["B1"] = f"{company} 기초재무분석"
    ws["B1"].font = _font(bold=True, size=14)
    ws["B2"] = "(단위: 원, 지표는 %·일·배)"
    ws["B2"].font = _font(italic=True)

    r = 4
    for title, rows, comments, ind_title, ind_rows, ind_comments in blocks:
        ws.cell(r, 2, title).font = _font(bold=True, size=12)
        r += 1
        for k, h in enumerate(["계정과목"] + [int(y) for y in years] + ["Level", "코멘트"], start=2):
            c = ws.cell(r, k, h)
            c.font = _font(bold=True)
            c.border = Border(bottom=Side(style="thin"))
        r += 1
        for row in rows:
            item = comments.get(row["label"])
            ws.cell(r, 2, row["label"]).font = _font(bold=row["level"] <= 1)
            ws.cell(r, 2).alignment = Alignment(indent=max(row["level"] - 1, 0))
            for k, col in enumerate(ycols):
                raw_col = get_column_letter(_FIRST_YEAR_COL + k)
                cell = ws[f"{col}{r}"]
                cell.value = f"=RAW!{raw_col}{row['raw_row']}"
                cell.number_format = _NUM_FMT
                cell.font = _font(bold=row["level"] <= 1)
            if item:
                ws.cell(r, level_col, f"Level {item.get('level', '')}").font = _font()
                c = ws.cell(r, comment_col, item.get("comment", ""))
                c.font = _font()
                c.fill = PatternFill("solid", fgColor=item.get("color_hex", "F2F2F2"))
                c.alignment = Alignment(wrap_text=True, vertical="center")
            else:
                # 핵심 계정이 아닌 행은 그룹으로 숨긴다(＋로 펼침) — 스킬 「엑셀 산출물 형식」 숨김 처리.
                ws.row_dimensions[r].outlineLevel = 1
                ws.row_dimensions[r].hidden = True
            r += 1

        r += 1
        ws.cell(r, 2, ind_title).font = _font(bold=True)
        r += 1
        for k, h in enumerate(["지표명"] + [int(y) for y in years] + ["Level", "코멘트"], start=2):
            c = ws.cell(r, k, h)
            c.font = _font(bold=True)
            c.border = Border(bottom=Side(style="thin"))
        r += 1
        for ind in ind_rows:
            item = ind_comments.get(ind["지표명"], {})
            ws.cell(r, 2, ind["지표명"]).font = _font()
            for k, col in enumerate(ycols):
                v = _num(str(ind.get(years[k], "")).replace(",", ""))
                cell = ws[f"{col}{r}"]
                cell.value = v if v is not None else "-"
                cell.number_format = "0.0"
                cell.font = _font()
                cell.alignment = Alignment(horizontal="right")
            if item:
                ws.cell(r, level_col, f"Level {item.get('level', '')}").font = _font()
                c = ws.cell(r, comment_col, item.get("comment", ""))
                c.font = _font()
                c.fill = PatternFill("solid", fgColor=item.get("color_hex", "F2F2F2"))
                c.alignment = Alignment(wrap_text=True, vertical="center")
            r += 1
        r += 2

    ws.cell(r, 2, "[분석 전제]").font = _font(bold=True)
    r += 1
    for line in notes:
        ws.cell(r, 2, f"- {line}").font = _font()
        r += 1
    r += 1
    ws.cell(r, 2, "[색상 범례]").font = _font(bold=True)
    r += 1
    for hex_color, label in _LEGEND:
        ws.cell(r, 2, "").fill = PatternFill("solid", fgColor=hex_color)
        ws.cell(r, 3, label).font = _font()
        r += 1

    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 30
    for col in ycols:
        ws.column_dimensions[col].width = 16
    ws.column_dimensions[get_column_letter(level_col)].width = 9
    ws.column_dimensions[get_column_letter(comment_col)].width = 52
    ws.freeze_panes = "C4"


# ── Graph 시트 ────────────────────────────────────────────────────────────────

def _graph_specs(bs_rows, is_rows):
    """스킬 Q4-4 표준 그래프 후보를 우선순위 순으로 만들고, 데이터가 있는 것 중 5개를 쓴다.
    각 spec: (제목, 구분 머리글, [(계열명, 값 수식 함수, 축)], 막대 묶음 방식, 축 단위)."""
    def find(rows, pattern, exclude=None, top_only=False):
        for row in rows:
            if top_only and row["level"] > 1:
                continue
            label = row["label"]
            if _RATIO_ROW_RE.search(label):
                continue
            if re.search(pattern, label) and not (exclude and re.search(exclude, label)):
                return row
        return None

    def won(row):
        return lambda col: f"=RAW!{col}{row['raw_row']}/100000000"

    def ratio(num, den, minus=None):
        if minus:
            return lambda col: (f"=IFERROR((RAW!{col}{num['raw_row']}-RAW!{col}{minus['raw_row']})"
                                f"/RAW!{col}{den['raw_row']}*100,\"\")")
        return lambda col: f"=IFERROR(RAW!{col}{num['raw_row']}/RAW!{col}{den['raw_row']}*100,\"\")"

    rev = find(is_rows, r"매출액|매출$|영업수익", top_only=True)
    cogs = find(is_rows, r"매출원가", top_only=True)
    sga = find(is_rows, r"판매비|판관비|관리비", top_only=True)
    op = find(is_rows, r"영업이익|영업손익|영업손실", top_only=True)
    ca = find(bs_rows, r"유동자산", r"비유동")
    cl = find(bs_rows, r"유동부채", r"비유동")
    inv = find(bs_rows, r"재고자산")
    cash = find(bs_rows, r"현금")
    liab = find(bs_rows, r"부채총계|부채합계", r"자본")
    equity = find(bs_rows, r"자본총계|자본합계", r"부채")
    short_b = find(bs_rows, r"단기차입금")
    long_b = find(bs_rows, r"장기차입금", r"유동성")
    ar = find(bs_rows, r"매출채권")
    ap = find(bs_rows, r"매입채무")

    specs = []
    if rev and (cogs or sga or op):
        series = []
        if sga:
            series.append(("판매관리비율", ratio(sga, rev), "left"))
        if op:
            series.append(("영업이익률", ratio(op, rev), "left"))
        if cogs:
            series.append(("매출원가율", ratio(cogs, rev), "right"))
        specs.append(("수익성 지표", "지표명", series, None, ("%", "%")))
    if liab and equity:
        specs.append(("자본 구조", "구분", [("부채총계", won(liab), "bar"), ("자본총계", won(equity), "bar"),
                                        ("부채비율", ratio(liab, equity), "right")], "clustered", ("억", "%")))
    if rev and rev.get("children") and len(rev["children"]) > 1:
        parts = [c for c in rev["children"] if any(v for v in c["values"])]
        if len(parts) > 1:
            specs.append(("매출 구성", "구분", [(re.sub(r"^\s*\d+\.\s*", "", c["label"]), won(c), "bar")
                                            for c in parts], "stacked", ("억", None)))
    if cash and ca and cl:
        series = [("현금성자산", won(cash), "bar"), ("유동비율", ratio(ca, cl), "right")]
        if inv:
            series.append(("당좌비율", ratio(ca, cl, minus=inv), "right"))
        specs.append(("현금·유동성", "구분", series, "clustered", ("억", "%")))
    if short_b or long_b:
        series = [(n, won(r), "bar") for n, r in (("단기차입금", short_b), ("장기차입금", long_b)) if r]
        specs.append(("차입금 구성", "구분", series, "stacked", ("억", None)))
    if ar or inv or ap:
        series = [(n, won(r), "bar") for n, r in (("매출채권", ar), ("재고자산", inv), ("매입채무", ap)) if r]
        specs.append(("운전자본 구성", "구분", series, "clustered", ("억", None)))
    return specs[:5]


def _axis_title(axis, text):
    axis.title = text
    axis.delete = False


def _value_labels():
    # 값만 표시한다(엑셀 기본값은 계열명·항목명까지 붙여 라벨이 겹친다).
    labels = DataLabelList()
    labels.showVal = True
    labels.showSerName = labels.showCatName = labels.showLegendKey = labels.showPercent = False
    return labels


def _straight_lines(chart):
    for series in chart.series:
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 6


def _write_graph(ws, years, specs):
    ycols = [3 + k for k in range(len(years))]
    ws["B1"] = "재무지표 그래프"
    ws["B1"].font = _font(bold=True, size=14)
    ws["B2"] = f"출처: 표준재무제표 {len(years)}개년 (금액 단위: 억, 비율 단위: %)"
    ws["B2"].font = _font(italic=True)
    block_h = 20
    for i, (title, head, series, grouping, units) in enumerate(specs):
        top = 4 + i * block_h
        ws.cell(top, 2, title).font = _font(bold=True, size=11)
        hdr = top + 1
        ws.cell(hdr, 2, head).font = _font(bold=True)
        for k, c in enumerate(ycols):
            ws.cell(hdr, c, int(years[k])).font = _font(bold=True)
        first = hdr + 1
        for j, (name, fn, _) in enumerate(series):
            rr = first + j
            ws.cell(rr, 2, name).font = _font()
            for k, c in enumerate(ycols):
                raw_col = get_column_letter(_FIRST_YEAR_COL + k)
                cell = ws.cell(rr, c, fn(raw_col))
                cell.number_format = "0.0"
                cell.font = _font()
        ws.cell(first + len(series), 2, f"출처: 표준재무제표 {len(years)}개년").font = _font(italic=True, size=8)

        cats = Reference(ws, min_col=ycols[0], max_col=ycols[-1], min_row=hdr)
        idx = {"bar": [], "left": [], "right": []}
        for j, (_, _, axis) in enumerate(series):
            idx[axis].append(first + j)

        def _ref(row):
            return Reference(ws, min_col=2, max_col=ycols[-1], min_row=row, max_row=row)

        if idx["bar"]:
            main = BarChart()
            main.type = "col"
            main.grouping = grouping or "clustered"
            if grouping == "stacked":
                main.overlap = 100
            for row in idx["bar"]:
                main.add_data(_ref(row), from_rows=True, titles_from_data=True)
            _axis_title(main.y_axis, f"(단위: {units[0]})")
        else:
            main = LineChart()
            for row in idx["left"]:
                main.add_data(_ref(row), from_rows=True, titles_from_data=True)
            _axis_title(main.y_axis, f"(단위: {units[0]})")
        main.set_categories(cats)
        main.x_axis.delete = False
        main.title = title  # 옆 원천 표 제목과 같은 글자 (Q4-4)
        main.dataLabels = _value_labels()
        main.x_axis.majorGridlines = None
        main.legend.position = "b"
        main.legend.overlay = False
        main.x_axis.tickLblPos = "low"  # 음수 값이 있어도 연도 라벨은 차트 맨 아래에
        if isinstance(main, LineChart):
            _straight_lines(main)

        if idx["right"]:
            sub = LineChart()
            for row in idx["right"]:
                sub.add_data(_ref(row), from_rows=True, titles_from_data=True)
            sub.y_axis.axId = 200
            sub.y_axis.crosses = "max"
            sub.y_axis.majorGridlines = None
            _axis_title(sub.y_axis, f"(단위: {units[1]})")
            sub.dataLabels = _value_labels()
            _straight_lines(sub)
            main += sub

        main.width, main.height = 17, 8.5
        ws.add_chart(main, f"{get_column_letter(ycols[-1] + 2)}{top}")

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 16
    for c in ycols:
        ws.column_dimensions[get_column_letter(c)].width = 10


# ── 진입점 ────────────────────────────────────────────────────────────────────

def build_analysis_workbook(company, df_bs, df_is, years, bs_comments, is_comments,
                            bs_indicator_rows, is_indicator_rows, bs_indicator_comments, is_indicator_comments):
    """Graph·Analysis·RAW 3시트 엑셀을 만들어 바이트로 돌려준다."""
    years = [str(y) for y in years]
    company = company or "진단기업"
    bs_rows = _statement_rows(df_bs, years)
    is_rows = _statement_rows(df_is, years)
    _assign_children(bs_rows, len(years), is_income=False)
    _assign_children(is_rows, len(years), is_income=True)

    wb = Workbook()
    graph = wb.active
    graph.title = "Graph"
    analysis = wb.create_sheet("Analysis")
    raw = wb.create_sheet("RAW")

    _write_raw(raw, company, years, bs_rows, is_rows)

    def by_label(items):
        return {i.get("account"): i for i in items or []}

    notes = [
        f"단위: 원(지표는 %·일·배). 출처: 고객 업로드 재무제표 {len(years)}개년({years[0]}~{years[-1]}), AI 추출 후 RAW 시트에 원본 그대로 기재",
        "활동성 지표(회수기간·회전일수·결제기간·현금전환주기)는 기말 잔액 기준. 공사미수금은 매출채권에 포함",
        "차입금 = 단기차입금 + 장기차입금 + 유동성장기부채 + 사채",
        "계정 선정·색상·코멘트는 기초재무진단 스킬 기준으로 AI가 자동 작성 — 검토 필요",
        "업종 평균 비교는 적용하지 않음(같은 기간 업종 자료 없음)",
    ]
    _write_analysis(analysis, company, years, [
        ("손익계산서", is_rows, by_label(is_comments), "손익계산서 관련 지표",
         is_indicator_rows or [], by_label(is_indicator_comments)),
        ("재무상태표", bs_rows, by_label(bs_comments), "재무상태표 관련 지표",
         bs_indicator_rows or [], by_label(bs_indicator_comments)),
    ], notes)
    _write_graph(graph, years, _graph_specs(bs_rows, is_rows))

    wb.calculation.fullCalcOnLoad = True  # 수식·차트 값은 엑셀이 열 때 계산한다
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
