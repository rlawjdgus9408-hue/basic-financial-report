"""
Section 4 (고객용): 최종보고서 생성

내부 제작용과 달리 사람이 지표를 고르거나 종합의견을 다듬는 과정이 없다 — 자료 업로드
직후 sections.auto_pipeline이 이미 지표 계산·그래프 생성·AI 종합의견 작성까지 끝내
session_state에 채워뒀으므로, 여기서는 그 결과를 고객용 템플릿에 그대로 채워 넣기만 한다.
보고서가 만들어지면 (SMTP 설정이 돼 있을 때) 내부 담당자에게 메일로도 전송할 수 있다.
"""
import hashlib
import io
import time
import datetime
from pathlib import Path

import re

import streamlit as st
from docxtpl import DocxTemplate
from docx.shared import Mm, Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from sections import email_notify
from sections import submission_log

_CUSTOMER_TEMPLATE = Path(__file__).parent.parent / "assets" / "templates" / "[그로스파이낸스]_고객용_템플릿.docx"

_FINANCE_TABLE_MARKER = "%%FINANCE_TABLE%%"
_ACCOUNT_TABLE_MARKER = "%%ACCOUNT_TABLE%%"

# 템플릿에 원래 있던 "◆ 재무지표 그래프" ~ "◆ 재무분석" 구간은 고정 5슬롯(차트 5개)짜리인데,
# 이제 중요도 기준 2~3개 가변 개수로 바뀌어서 그 구간 전체를 통째로 다시 만든다. 이 두
# 제목은 jinja 변수가 아니라 템플릿에 원래 박혀있는 고유한 정적 텍스트라, sentinel 없이
# 그 글자 자체를 anchor로 쓴다.
_CHART_SECTION_START = "◆ 재무지표 그래프"
_CHART_SECTION_END = "◆ 재무분석"

# 엑셀에서 복붙한 템플릿이라 run마다 폰트가 제각각인데, 전체를 이 폰트로 고정한다.
_REPORT_FONT = "맑은 고딕"


def _shade_cell(cell, hex_color):
    """표 셀 배경색을 입힌다(python-docx 표준 기법 — w:tcPr/w:shd를 직접 추가)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


_ROMAN_NUMERAL_RE = re.compile(r"[Ⅰ-Ⅿ]")  # Ⅰ Ⅱ Ⅲ … — 재무상태표 대분류 표기


def _bold_cell(cell):
    for run in cell.paragraphs[0].runs:
        run.bold = True


def _clear_cell_borders(cell):
    """셀 테두리를 없앤다(w:tcBorders의 네 변을 전부 nil로) — python-docx table style은
    열 단위 제어가 안 돼 셀 XML을 직접 만져야 한다."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _set_column_width(table, col_index, width_cm):
    """해당 열의 Column.width와 그 열에 속한 모든 셀의 width를 함께 지정한다 — python-docx는
    Column.width만 설정하면 Word에서 무시되는 경우가 있어(셀 폭이 우선 적용됨) 셀에도
    같은 폭을 다시 지정해야 실제로 반영된다."""
    table.columns[col_index].width = Cm(width_cm)
    for row in table.rows:
        row.cells[col_index].width = Cm(width_cm)


def _new_table(doc, header_row, data_rows, row_colors=None, bold_first_cell_rows=None,
               strip_border_col=None, col_widths_cm=None):
    """header_row/data_rows로 Word 표를 만들어 반환한다(아직 문서 끝에 붙어있는 상태 —
    호출부가 원하는 위치로 옮겨야 한다). 헤더 행은 항상 굵게 표시한다.
    - row_colors[i]가 있으면 i번째 데이터 행 전체에 그 hex 배경색을 입힌다.
    - bold_first_cell_rows[i]가 True면 그 데이터 행의 첫 번째 셀(계정명 등)을 굵게 표시한다.
    - strip_border_col이 주어지면 그 열(헤더 포함) 전체의 테두리를 없앤다.
    - col_widths_cm이 주어지면(각 열의 cm 폭 리스트) 표를 고정폭으로 만든다."""
    n_cols = len(header_row)
    table = doc.docx.add_table(rows=1 + len(data_rows), cols=n_cols)
    table.style = "Table Grid"

    for j, h in enumerate(header_row):
        cell = table.rows[0].cells[j]
        cell.text = str(h)
        _bold_cell(cell)

    for i, row_vals in enumerate(data_rows):
        row = table.rows[i + 1]
        for j, v in enumerate(row_vals):
            row.cells[j].text = str(v)
        if bold_first_cell_rows and bold_first_cell_rows[i]:
            _bold_cell(row.cells[0])
        color = row_colors[i] if row_colors else None
        if color:
            for cell in row.cells:
                _shade_cell(cell, color)

    if strip_border_col is not None:
        for row in table.rows:
            _clear_cell_borders(row.cells[strip_border_col])

    if col_widths_cm:
        table.autofit = False
        for j, width in enumerate(col_widths_cm):
            if width is not None:
                _set_column_width(table, j, width)

    return table


def _commented_table_col_widths(n_year_cols):
    """계정/지표 코멘트 표의 열 폭 — 라벨·연도 열은 좁게, 코멘트 열만 넓게 고정한다."""
    return [2.5] + [2.0] * n_year_cols + [6.0]


def _new_heading_paragraph(doc, text):
    """굵게 표시되는 소제목 단락을 만들어 반환한다(표와 마찬가지로 문서 끝에 붙은 상태)."""
    p = doc.docx.add_paragraph()
    p.text = text
    for run in p.runs:
        run.bold = True
    return p


def _insert_blocks_after_marker(doc, marker_text, blocks):
    """marker_text가 들어있는 단락을 찾아 비우고, 그 바로 뒤에 blocks(Paragraph._p 또는
    Table._tbl 엘리먼트 리스트)를 순서대로 이어붙인다.

    python-docx에는 '특정 위치에 삽입' API가 없어, 각 블록을 일단 문서 끝에 만든 뒤
    XML 엘리먼트를 원하는 자리로 옮기는 표준 기법을 쓴다(addnext가 기존 부모에서 자동으로
    떼어내 옮겨준다). marker가 템플릿에 없으면(다른 템플릿을 쓰는 경우 등) 조용히
    건너뛴다 — 보고서 생성 자체를 막지 않기 위함."""
    anchor = next((p for p in doc.docx.paragraphs if marker_text in p.text), None)
    if anchor is None:
        return
    anchor.text = ""

    last = anchor._p
    for block in blocks:
        last.addnext(block)
        last = block


def _format_won_raw(value):
    """원 단위 그대로, 천단위 콤마만 넣는다 — 재무상태표/손익계산서는 참고 샘플(㈜초이인터
    네셔널 보고서)처럼 "(단위: 원)" 원본 금액으로 보여준다(억 변환은 차트 옆 표에서만 쓴다)."""
    try:
        return f"{int(round(float(value))):,}"
    except (TypeError, ValueError):
        return str(value)


def _account_row(item, df, years):
    """계정별 코멘트 항목 하나를 표의 한 행([계정과목, 연도별 값..., 코멘트])으로 만든다.
    df(customer_df_bs/is)에서 같은 계정과목 행을 찾아 연도별 값을 채우고, 못 찾으면
    (AI가 계정명을 원본과 다르게 적은 경우 등) '-'로 비워둔다 — 룩업 실패가 코멘트
    표시 자체를 막지 않게 한다."""
    account = item.get("account", "")
    values = ["-"] * len(years)
    if df is not None:
        match = df[df["계정과목"] == account]
        if not match.empty:
            row = match.iloc[0]
            values = [_format_won_raw(row.get(y, "-")) for y in years]
    return [account] + values + [item.get("comment", "")]


def _indicator_row(item, indicator_rows, years):
    """지표 코멘트 항목 하나를 표의 한 행([지표명, 연도별 값..., 코멘트])으로 만든다.
    indicator_rows(selected_indicators_bs_data/is_data)에서 같은 지표명 행을 찾아 이미
    포맷된 값을 그대로 쓴다(지표는 비율이라 _account_row와 달리 원/억 변환이 필요 없다)."""
    name = item.get("account", "")
    values = ["-"] * len(years)
    match = next((r for r in (indicator_rows or []) if r.get("지표명") == name), None)
    if match:
        values = [match.get(y, "-") for y in years]
    return [name] + values + [item.get("comment", "")]


def _new_chart_block(doc, spec, years_count):
    """차트 하나를 소제목 + (왼쪽 차트 이미지 / 오른쪽 원본 데이터 표) 1행2열 표 + 출처
    캡션으로 구성한다. 차트 제목·범례는 auto_pipeline.py가 spec['table_header']/
    ['table_rows']에 이미 같은 문자열로 넣어뒀으므로 여기서는 그대로 옮겨 쓰기만 한다."""
    heading = _new_heading_paragraph(doc, spec["title"])

    outer = doc.docx.add_table(rows=1, cols=2)
    outer.style = "Table Grid"
    left_cell, right_cell = outer.rows[0].cells
    left_cell.paragraphs[0].add_run().add_picture(io.BytesIO(spec["png_bytes"]), width=Mm(80))

    header = spec["table_header"]
    rows = spec["table_rows"]
    nested = right_cell.add_table(rows=1 + len(rows), cols=len(header))
    nested.style = "Table Grid"
    for j, h in enumerate(header):
        cell = nested.rows[0].cells[j]
        cell.text = str(h)
        _bold_cell(cell)
    for i, row_vals in enumerate(rows):
        for j, v in enumerate(row_vals):
            nested.rows[i + 1].cells[j].text = str(v)

    caption = doc.docx.add_paragraph()
    caption.text = f"출처: 표준재무제표 {years_count}개년"
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8)

    return [heading._p, outer._tbl, caption._p]


def _rebuild_chart_section(doc, chart_specs, years_count):
    """"◆ 재무지표 그래프" ~ "◆ 재무분석" 사이(옛 5슬롯 소제목·이미지 자리)를 전부 지우고,
    그 자리에 chart_specs(2~3개)로 만든 블록을 끼워 넣는다. 두 anchor 중 하나라도 못 찾으면
    (다른 템플릿을 쓰는 경우 등) 조용히 건너뛴다 — 보고서 생성 자체를 막지 않기 위함."""
    start = next((p for p in doc.docx.paragraphs if p.text.strip() == _CHART_SECTION_START), None)
    end = next((p for p in doc.docx.paragraphs if p.text.strip() == _CHART_SECTION_END), None)
    if start is None or end is None:
        return

    el = start._p.getnext()
    while el is not None and el is not end._p:
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt

    last = start._p
    for spec in chart_specs:
        for block in _new_chart_block(doc, spec, years_count):
            last.addnext(block)
            last = block


_HEADER_RELTYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header"
_FOOTER_RELTYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer"


def _set_run_font(run_element, font_name):
    rPr = run_element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run_element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font_name)


def _force_font(doc, font_name=_REPORT_FONT):
    """문서 전체(본문·표·중첩표·머리글·바닥글)의 글꼴을 강제로 통일한다. doc은 DocxTemplate
    인스턴스(본문은 doc.docx, 헤더/푸터는 doc.docx._part.rels로 접근) — run.font.name만
    설정하면 영문만 바뀌고 한글 표시 폰트는 그대로인 python-docx의 알려진 함정이 있어,
    w:rFonts의 w:eastAsia까지 직접 XML로 설정해야 한다.

    헤더/푸터는 일부러 python-docx의 Section.header/footer가 아니라 doc.docx._part.rels로
    파트를 직접 찾아 그 XML 엘리먼트를 수정한다 — docxtpl이 render() 중에 헤더/푸터 파트를
    통째로 새 XmlPart로 교체하는데(map_headers_footers_xml), Section.header를 통해 수정하면
    바뀐 내용이 save() 후 사라지는 현상이 있었다(다른 경로로 만들어진 파트 참조라 그런 것으로
    보임) — 방금 교체된 그 파트 객체를 relKey로 직접 집어 수정하면 확실히 저장된다."""
    def _apply(paragraphs):
        for p in paragraphs:
            for run in p.runs:
                run.font.name = font_name
                _set_run_font(run._element, font_name)

    def _apply_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    _apply(cell.paragraphs)
                    _apply_tables(cell.tables)

    docx = doc.docx
    _apply(docx.paragraphs)
    _apply_tables(docx.tables)

    for rel in docx._part.rels.values():
        if rel.reltype in (_HEADER_RELTYPE, _FOOTER_RELTYPE):
            for run_element in rel.target_part.element.iter(qn("w:r")):
                _set_run_font(run_element, font_name)


def render_report_generation(**_):
    """최종보고서 생성 렌더링 (콘텐츠 영역 — 실제 생성/다운로드는 사이드바에서)"""
    st.markdown("---")
    st.markdown("### 4. 최종보고서 생성")
    st.info("사이드바의 '보고서 생성' 버튼을 클릭하여 보고서를 생성하세요.")
    st.caption(
        "※ 본 보고서는 AI가 자동으로 생성한 결과로, 실제와 다르거나 정확하지 않을 수 있습니다. "
        "정확한 재무 분석이 필요하시면 그로스파이낸스에 문의해 주세요."
    )

    pipeline_error = st.session_state.get('auto_pipeline_error')
    if pipeline_error:
        st.warning(f"자동분석 중 일부 항목에 문제가 있었습니다: {pipeline_error}")


def _build_report_bytes(company_info, template_path, check_results, score,
                         selected_dirs, dir_etc, selected_mats, exec_summary):
    doc = DocxTemplate(template_path)

    context = {
        'today': datetime.datetime.now().strftime("%Y. %m."),
        'company_name': company_info.get('company_name', ''),
        'biz_type': company_info.get('biz_type', ''),
        'ceo_name': company_info.get('ceo_name', ''),
        'biz_start_date': company_info.get('biz_start_date', ''),
        'biz_no': company_info.get('biz_no', ''),
        'phone': company_info.get('phone', ''),
        'email': company_info.get('email', ''),
        'address': company_info.get('address', ''),
        'emp_count': company_info.get('emp_count', ''),
        'erp_system': ', '.join(company_info.get('erp_system', [])) if isinstance(company_info.get('erp_system'), list) else company_info.get('erp_system', ''),
        'special_note': company_info.get('special_note', ''),
        'exec_summary': exec_summary,
        'finance_comment': _FINANCE_TABLE_MARKER,
        'account_comment': _ACCOUNT_TABLE_MARKER,
        'score': score,
    }

    for i, res in enumerate(check_results):
        num = i + 1
        is_yes = (res == "예")
        context[f'r{num}y'] = "■" if is_yes else "□"
        context[f'r{num}n'] = "■" if not is_yes else "□"
        context[f's{num}y'] = "■" if is_yes else "□"
        context[f's{num}n'] = "■" if not is_yes else "□"

    for i, opt in enumerate(["안정적 성장", "투자 유치 (사업확장)", "IPO/M&A 등 Exit"]):
        context[f'd{i+1}'] = "■" if opt in selected_dirs else "□"

    context['d_etc'] = "■" if dir_etc else "□"
    context['d_etc_val'] = dir_etc if dir_etc else "          "

    for i, opt in enumerate(["사업자등록증", "재무제표", "회사소개서", "기타 양식"]):
        context[f'm{i+1}'] = "■" if opt in selected_mats else "□"

    # 템플릿에 원래 있던 5개 고정 차트 슬롯은 더 이상 쓰지 않는다(아래 _rebuild_chart_section이
    # 그 구간을 통째로 지우고 2~3개 가변 블록으로 다시 만든다) — render() 단계에서는 빈
    # 문자열만 채워 undefined 변수 오류 없이 넘어가게 한다.
    for var in ('profitability_chart', 'stability_chart', 'activity_chart', 'growth_chart', 'roi_chart'):
        context[var] = ''

    doc.render(context)

    # finance_comment(지표 표)·account_comment(계정별 코멘트 표)는 jinja 변수 하나로 끝낼 수
    # 없는 실제 Word 표라서, render() 전에는 자리표시 sentinel만 넣어두고 render() 후에
    # 그 단락을 찾아 표로 바꿔치기한다(위 _insert_blocks_after_marker 참고).
    # finance_comment는 더 이상 따로 쓰지 않는다 — 재무상태표/손익계산서 관련 지표 표를
    # 각자 해당 계정 표 바로 밑에 넣도록 account_comment 한 자리에 전부 모아 끼워 넣는다
    # (둘 다 내가 만든 sentinel이라 원래 템플릿의 두 자리 순서에 얽매일 필요가 없다).
    _insert_blocks_after_marker(doc, _FINANCE_TABLE_MARKER, [])

    years = st.session_state.get('years', [])
    col_widths = _commented_table_col_widths(len(years))
    bs_df = st.session_state.get('customer_df_bs')
    is_df = st.session_state.get('customer_df_is')
    bs_accounts = st.session_state.get('account_bs_comments') or []
    is_accounts = st.session_state.get('account_is_comments') or []
    bs_indicators = st.session_state.get('indicator_bs_comments') or []
    is_indicators = st.session_state.get('indicator_is_comments') or []
    bs_indicator_rows = st.session_state.get('selected_indicators_bs_data') or []
    is_indicator_rows = st.session_state.get('selected_indicators_is_data') or []

    account_blocks = []
    if bs_accounts:
        header = ['계정과목'] + list(years) + ['코멘트']
        rows = [_account_row(item, bs_df, years) for item in bs_accounts]
        colors = [item.get('color_hex') for item in bs_accounts]
        bold_rows = [bool(_ROMAN_NUMERAL_RE.search(item.get('account', ''))) for item in bs_accounts]
        account_blocks.append(_new_heading_paragraph(doc, "■ 재무상태표")._p)
        account_blocks.append(_new_table(
            doc, header, rows, colors, bold_first_cell_rows=bold_rows,
            strip_border_col=len(header) - 1, col_widths_cm=col_widths,
        )._tbl)
    if bs_indicators:
        header = ['지표명'] + list(years) + ['코멘트']
        rows = [_indicator_row(item, bs_indicator_rows, years) for item in bs_indicators]
        colors = [item.get('color_hex') for item in bs_indicators]
        account_blocks.append(_new_heading_paragraph(doc, "■ 재무상태표 관련 지표")._p)
        account_blocks.append(_new_table(
            doc, header, rows, colors, strip_border_col=len(header) - 1, col_widths_cm=col_widths,
        )._tbl)
    if is_accounts:
        header = ['계정과목'] + list(years) + ['코멘트']
        rows = [_account_row(item, is_df, years) for item in is_accounts]
        colors = [item.get('color_hex') for item in is_accounts]
        bold_rows = [bool(_ROMAN_NUMERAL_RE.search(item.get('account', ''))) for item in is_accounts]
        account_blocks.append(_new_heading_paragraph(doc, "■ 손익계산서")._p)
        account_blocks.append(_new_table(
            doc, header, rows, colors, bold_first_cell_rows=bold_rows,
            strip_border_col=len(header) - 1, col_widths_cm=col_widths,
        )._tbl)
    if is_indicators:
        header = ['지표명'] + list(years) + ['코멘트']
        rows = [_indicator_row(item, is_indicator_rows, years) for item in is_indicators]
        colors = [item.get('color_hex') for item in is_indicators]
        account_blocks.append(_new_heading_paragraph(doc, "■ 손익계산서 관련 지표")._p)
        account_blocks.append(_new_table(
            doc, header, rows, colors, strip_border_col=len(header) - 1, col_widths_cm=col_widths,
        )._tbl)
    _insert_blocks_after_marker(doc, _ACCOUNT_TABLE_MARKER, account_blocks)

    chart_specs = st.session_state.get('auto_chart_specs') or []
    _rebuild_chart_section(doc, chart_specs, len(years))

    _force_font(doc)

    # 화면에는 "AI 자동 생성" 안내 문구만 보이고 실제 파일에는 아무 표시가 안 남아있었다 —
    # 파일만 외부로 유출/재배포되는 경우에도 출처를 추적할 수 있도록, 눈에 보이진 않지만
    # 문서 속성(파일 정보)에 출처·생성시각·문서ID를 남긴다. 레이아웃에 영향 없음.
    generated_at = datetime.datetime.now()
    doc_id = hashlib.sha1(
        f"{company_info.get('company_name', '')}|{generated_at.isoformat()}".encode("utf-8")
    ).hexdigest()[:12]
    doc.docx.core_properties.author = "GrowthFinance(그로스파이낸스)"
    doc.docx.core_properties.comments = (
        f"그로스파이낸스 기초재무진단 자동분석 시스템 생성 문서 · "
        f"생성일시 {generated_at.strftime('%Y-%m-%d %H:%M:%S')} · 문서ID {doc_id}"
    )

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_report_sidebar(company_info, template_file, check_results, score,
                           selected_dirs, dir_etc, selected_mats, exec_summary):
    """사이드바의 '보고서 생성' 버튼 — 실제 docx 렌더링, 다운로드, (설정 시) 내부 메일 전송."""
    st.sidebar.markdown("---")
    st.sidebar.subheader("최종보고서")

    template_path = str(_CUSTOMER_TEMPLATE) if _CUSTOMER_TEMPLATE.exists() else template_file
    if not template_path:
        st.sidebar.error("보고서 템플릿 파일을 찾을 수 없습니다.")
        return

    if st.sidebar.button("보고서 생성", use_container_width=True):
        progress_text = "데이터를 분석하여 보고서를 생성 중입니다. 잠시만 기다려주세요..."
        my_bar = st.sidebar.progress(0, text=progress_text)

        for percent_complete in range(100):
            time.sleep(0.015)
            my_bar.progress(percent_complete + 1, text=progress_text)

        my_bar.empty()

        # 고객용 흐름은 업로드와 보고서 생성 사이에 사람이 데이터를 검토하는 화면이 없어,
        # 예상과 다른 추출 결과나 템플릿 변수 불일치가 그대로 보고서 생성 단계에서 터질 수
        # 있다 — 여기서 막지 않으면 고객 화면에 날것의 에러 트레이스백이 그대로 노출된다.
        try:
            report_bytes = _build_report_bytes(
                company_info, template_path, check_results, score,
                selected_dirs, dir_etc, selected_mats, exec_summary,
            )
        except Exception as error:
            st.sidebar.error(
                f"보고서 생성 중 오류가 발생했습니다: {error}\n"
                "잠시 후 다시 시도하거나, 계속되면 그로스파이낸스에 문의해 주세요."
            )
        else:
            file_name_prefix = company_info.get('company_name', '') or "진단기업"
            st.session_state['report_bytes'] = report_bytes
            created_date = datetime.datetime.now().strftime('%Y%m%d')
            st.session_state['report_filename'] = f"{file_name_prefix}_재무진단보고서_{created_date}.docx"
            st.session_state['report_company'] = file_name_prefix
            st.session_state.pop('report_email_sent', None)
            st.sidebar.success("보고서 생성 완료!")
            submission_log.log_event(
                "보고서 생성", company_info.get('company_name', ''),
                company_info.get('phone', ''), company_info.get('email', ''),
            )

    auto_es_warnings = st.session_state.get('auto_es_warnings')
    if auto_es_warnings:
        st.sidebar.warning(
            "자동 생성된 종합의견에 표기 규칙 위반이 있어 확인이 필요합니다:\n"
            + "\n".join(f"- {w}" for w in auto_es_warnings)
        )

    auto_account_comment_warnings = st.session_state.get('auto_account_comment_warnings')
    if auto_account_comment_warnings:
        st.sidebar.warning(
            "자동 생성된 계정별 코멘트에 표기 규칙 위반이 있어 확인이 필요합니다:\n"
            + "\n".join(f"- {w}" for w in auto_account_comment_warnings)
        )

    if st.session_state.get('report_bytes'):
        file_name_prefix = st.session_state.get('report_company', '진단기업')
        st.sidebar.download_button(
            f"{file_name_prefix} 보고서 다운로드",
            data=st.session_state['report_bytes'],
            file_name=st.session_state['report_filename'],
            use_container_width=True
        )
        st.sidebar.caption(
            "※ AI 자동 생성 결과로 부정확할 수 있습니다. 정확한 분석은 그로스파이낸스에 문의해 주세요."
        )

        if email_notify.is_configured():
            if st.sidebar.button("이사님께 결과 전송", use_container_width=True, key="btn_send_report_email"):
                try:
                    email_notify.send_report_notification(
                        company_name=company_info.get('company_name', ''),
                        contact=company_info.get('phone', ''),
                        report_bytes=st.session_state['report_bytes'],
                        report_filename=st.session_state['report_filename'],
                    )
                    st.session_state['report_email_sent'] = True
                    st.sidebar.success(f"{', '.join(email_notify.RECIPIENTS)}에 메일을 보냈습니다.")
                except Exception as error:
                    st.sidebar.error(f"메일 전송에 실패했습니다: {error}")
            elif st.session_state.get('report_email_sent'):
                st.sidebar.caption("✓ 이미 전송했습니다.")
        else:
            st.sidebar.caption("※ 메일 발송 계정이 아직 설정되지 않아 '결과 전송' 버튼은 숨겨져 있습니다.")
